import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from services.retrieval.extract import InvalidDocument, extract, split_text


class ExtractionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_markdown_unicode_and_stable_chunks(self) -> None:
        path = self.root / "source.md"
        text = "Français العربية 中文\n" * 500
        path.write_text(text)
        self.assertEqual(extract(path), text.strip())
        chunks = split_text(text)
        self.assertEqual(chunks, split_text(text))
        self.assertTrue(all(0 < len(c) <= 2400 for c in chunks))
        self.assertIn("العربية", chunks[0])
        self.assertEqual(split_text(""), [])

    def test_html_excludes_active_content(self) -> None:
        path = self.root / "page.html"
        path.write_text(
            "<p>A &amp; B</p><script>bad</script><style>bad</style><p>中文</p>"
        )
        self.assertEqual(extract(path), "A & B\n中文")

    def test_docx_preserves_paragraph_table_order(self) -> None:
        path = self.root / "document.docx"
        doc = Document()
        doc.add_paragraph("avant العربية")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "中文"
        doc.add_paragraph("après")
        doc.save(str(path))
        self.assertEqual(extract(path), "avant العربية\n中文\naprès")

    def test_pdf_text_and_blank_page(self) -> None:
        path = self.root / "document.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 700 Td (Evidence PDF) Tj ET")
        page[NameObject("/Contents")] = content
        writer.add_blank_page(width=600, height=800)
        writer.write(path)
        self.assertEqual(extract(path), "Evidence PDF")

    def test_rejects_corruption_empty_and_unsupported(self) -> None:
        for suffix, content in [
            ("pdf", b"broken"),
            ("docx", b"broken"),
            ("md", b""),
            ("html", b"<script>x</script>"),
            ("txt", b"text"),
            ("md", b"\xff"),
        ]:
            with self.subTest(suffix=suffix, content=content):
                path = self.root / ("bad." + suffix)
                path.write_bytes(content)
                with self.assertRaises(InvalidDocument):
                    extract(path)

    def test_rejects_limits_and_symlinks(self) -> None:
        path = self.root / "large.md"
        path.write_bytes(b"x" * (10 * 1024 * 1024 + 1))
        with self.assertRaises(InvalidDocument):
            extract(path)
        path.write_text("x" * 1_000_001)
        with self.assertRaises(InvalidDocument):
            extract(path)
        link = self.root / "link.md"
        link.symlink_to(path)
        with self.assertRaises(InvalidDocument):
            extract(link)
        archive = self.root / "large.docx"
        with ZipFile(archive, "w") as z:
            z.writestr("word/document.xml", b"x" * 1_000_001)
        with self.assertRaises(InvalidDocument):
            extract(archive)

    def test_chunk_overlap_and_input_unchanged(self) -> None:
        text = "".join(str(i % 10) for i in range(7000))
        before = hashlib.sha256(text.encode()).hexdigest()
        chunks = split_text(text)
        self.assertEqual(chunks[0][-200:], chunks[1][:200])
        self.assertEqual(chunks[-1][-100:], text[-100:])
        self.assertEqual(hashlib.sha256(text.encode()).hexdigest(), before)

    def test_cli_and_encrypted_pdf(self) -> None:
        path = self.root / "cli.md"
        path.write_text("preuve العربية 中文")
        result = subprocess.run(
            [sys.executable, "-m", "services.retrieval.extract", str(path)],
            capture_output=True,
            text=True,
            timeout=20,
            check=True,
        )
        self.assertEqual(json.loads(result.stdout), ["preuve العربية 中文"])
        self.assertIn('"event": "extracted"', result.stderr)
        writer = PdfWriter()
        writer.add_blank_page(width=600, height=800)
        writer.encrypt("test-only-password")
        pdf = self.root / "encrypted.pdf"
        writer.write(pdf)
        with self.assertRaisesRegex(InvalidDocument, "encrypted_pdf"):
            extract(pdf)

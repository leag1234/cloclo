"""Offline acceptance tests for the Task H renderer."""

import json
import os
import tempfile
import unittest
from pathlib import Path
from typing import cast
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pydantic import ValidationError
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject
from pypdf.errors import PdfReadError

from packages.doc_render import Content, Section, VerificationError, render, verify
from packages.doc_render.layout import HEIGHT, MARGIN, pages


class DocumentRenderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.environment = patch.dict(os.environ, {"ATLAS_SEARCH_PROVIDER": "serpapi"})
        self.environment.start()
        self.addCleanup(self.environment.stop)
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        Image.new("RGB", (100, 50), (20, 90, 160)).save(self.root / "figure.png")
        self.fixture = json.loads(
            (Path(__file__).parent / "fixtures/doc_render.json").read_text()
        )
        self.content = Content.model_validate(self.fixture)

    def test_formats_reopen_preserve_content_and_are_deterministic(self) -> None:
        for suffix, count in (("pptx", 3), ("pdf", 3), ("docx", 1), ("xlsx", 2)):
            with self.subTest(format=suffix):
                first = self.root / f"first.{suffix}"
                result = render(self.content, first, image_root=self.root)
                self.assertEqual(result.count, count)
                self.assertTrue(all(value.strip() for value in result.text))
                self.assertEqual(result.pagination_verified, suffix in {"pdf", "pptx"})
                second = self.root / f"second.{suffix}"
                with patch("time.time", return_value=2_000_000_000):
                    render(self.content, second, image_root=self.root)
                self.assertEqual(first.read_bytes(), second.read_bytes())
                if suffix != "pdf":
                    with ZipFile(first) as archive:
                        self.assertIsNone(archive.testzip())
                verify(first, expected=self.content, expected_count=count)

    def test_content_only_and_literal_spreadsheet_cells(self) -> None:
        content = Content(
            title="Provided title",
            sections=(
                Section(
                    heading="Provided heading",
                    paragraphs=("=1+1", "@SUM(A1)", "+42", "-42"),
                ),
            ),
        )
        for suffix in ("pdf", "pptx", "docx", "xlsx"):
            path = self.root / f"literal.{suffix}"
            result = render(content, path, image_root=self.root)
            self.assertEqual(
                "".join("".join(result.text).split()),
                "".join("".join(content.texts()).split()),
            )
        workbook = load_workbook(self.root / "literal.xlsx", data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.worksheets[0]["A3"].value, "=1+1")
        self.assertTrue(
            all(cell.data_type == "s" for row in workbook.worksheets[0] for cell in row)
        )

    def test_heading_prefixes_preserve_all_sections(self) -> None:
        content = Content(
            title="Provided title",
            sections=tuple(
                Section(heading=f"Section {index}", paragraphs=(f"Body {index}",))
                for index in range(1, 11)
            ),
        )
        for suffix in ("pdf", "pptx"):
            path = self.root / f"prefixes.{suffix}"
            render(content, path, image_root=self.root)
            result = verify(path, expected=content, expected_count=10)
            self.assertIn("Section 10", result.text[-1])

    def test_schema_rejects_empty_ragged_and_unknown_content(self) -> None:
        for data in (
            {"title": "x", "sections": []},
            {"title": "x", "sections": [{"heading": "x"}]},
            {"title": " ", "sections": [{"heading": "x", "bullets": ["x"]}]},
            {"title": "x", "sections": [{"heading": "x", "bullets": ["\x00"]}]},
            {
                "title": "x",
                "sections": [
                    {
                        "heading": "x",
                        "tables": [{"headers": ["a", "b"], "rows": [["a"]]}],
                    }
                ],
            },
            {**self.fixture, "invented": True},
        ):
            with self.subTest(data=data), self.assertRaises(ValidationError):
                Content.model_validate(data)

    def test_image_boundaries_and_atomic_failure(self) -> None:
        destination = self.root / "existing.pdf"
        destination.write_bytes(b"original")
        (self.root / "escape.png").symlink_to(Path(__file__).resolve())
        for value in ("missing.png", "escape.png", str(self.root / "figure.png")):
            data = json.loads(json.dumps(self.fixture))
            data["sections"][0]["images"] = [value]
            with (
                self.subTest(path=value),
                self.assertRaises((ValueError, FileNotFoundError)),
            ):
                render(Content.model_validate(data), destination, image_root=self.root)
            self.assertEqual(destination.read_bytes(), b"original")

    def test_pagination_and_no_silent_truncation(self) -> None:
        content = Content(
            title="Long content",
            sections=(
                Section(
                    heading="Source heading",
                    paragraphs=tuple(
                        f"Paragraph {i}: " + "longword " * 60 for i in range(18)
                    ),
                ),
            ),
        )
        layout = pages(content, self.root)
        self.assertGreater(len(layout), 2)
        for page in layout:
            self.assertTrue(
                all(item.top + item.block.height <= HEIGHT - MARGIN for item in page)
            )
        for suffix in ("pdf", "pptx"):
            path = self.root / f"long.{suffix}"
            render(content, path, image_root=self.root)
            verify(
                path,
                expected=content,
                expected_count=len(layout),
                require_pagination=True,
            )

    def test_verification_rejects_empty_corrupt_missing_and_wrong_count(self) -> None:
        path = self.root / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=720, height=405)
        writer.write(path)
        with self.assertRaisesRegex(VerificationError, "empty"):
            verify(path)
        path.write_bytes(b"invalid")
        with self.assertRaises(PdfReadError):
            verify(path)
        path = self.root / "blank.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(str(path))
        with self.assertRaisesRegex(VerificationError, "empty"):
            verify(path)
        path = self.root / "valid.pdf"
        render(self.content, path, image_root=self.root)
        with self.assertRaisesRegex(VerificationError, "unit_count_mismatch"):
            verify(path, expected_count=99)
        missing = Content(title="Absent source text", sections=self.content.sections)
        with self.assertRaisesRegex(VerificationError, "missing_expected_text"):
            verify(path, expected=missing)
        for suffix in ("docx", "xlsx"):
            path = self.root / f"office.{suffix}"
            render(self.content, path, image_root=self.root)
            with self.assertRaisesRegex(
                VerificationError, "office_layout_engine_required"
            ):
                verify(path, require_pagination=True)

    def test_images_and_pdf_font_are_embedded(self) -> None:
        for suffix in ("pdf", "pptx", "docx", "xlsx"):
            render(self.content, self.root / f"images.{suffix}", image_root=self.root)
        reader = PdfReader(self.root / "images.pdf")
        self.assertEqual(sum(len(page.images) for page in reader.pages), 1)
        embedded = []
        for page in reader.pages:
            resources = cast(DictionaryObject, page["/Resources"])
            fonts = cast(DictionaryObject, resources["/Font"])
            for reference in fonts.values():
                font = cast(DictionaryObject, reference.get_object())
                if "/FontDescriptor" in font:
                    descriptor = cast(DictionaryObject, font["/FontDescriptor"])
                    embedded.append("/FontFile2" in descriptor)
        self.assertTrue(embedded)
        self.assertTrue(all(embedded))
        document = Document(str(self.root / "images.docx"))
        self.assertEqual(len(document.inline_shapes), 1)
        for suffix in ("pptx", "xlsx"):
            with ZipFile(self.root / f"images.{suffix}") as archive:
                self.assertEqual(
                    sum("/media/" in name for name in archive.namelist()), 1
                )

    def test_unsupported_glyph_and_huge_heading_fail(self) -> None:
        for title in (chr(0x1F680), "x" * 4000):
            content = Content(title=title, sections=self.content.sections)
            with self.assertRaises(ValueError):
                render(content, self.root / "reject.pdf", image_root=self.root)

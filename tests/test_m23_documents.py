"""Complete attachment extraction and HTTP routing regressions, without paid I/O."""

import asyncio
import base64
from collections.abc import AsyncGenerator
from decimal import Decimal
import io
import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import AsyncMock, patch
from typing import Any

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from services.orchestrator.chat_api import app
from services.orchestrator.chat_schema import ChatRequest
from services.orchestrator.chat_stream import response
from services.orchestrator.documents import Attachment, DocumentError, extract_document
from services.orchestrator.interactions import Interaction
from services.orchestrator.loop import Turn
from services.orchestrator.model import Configuration, GatewayModel
from services.orchestrator.stream_client import sink_context


def pdf(pages: int = 60) -> bytes:
    writer = PdfWriter()
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    for number in range(1, pages + 1):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        fact = (
            "Initial budget is 730 euros."
            if number == 1
            else "Midpoint review requires 19 samples."
            if number == 30
            else "Final recommendation is to postpone launch."
            if number == 60
            else "Every batch needs documented quality review."
        )
        lines = [f"Section {number}. {fact}"] + [
            f"Procedure {number}.{step}: record inspection results and retain signed evidence."
            for step in range(1, 41)
        ]
        stream.set_data(
            (
                "BT /F1 10 Tf 50 740 Td 15 TL "
                + " ".join(f"({line}) Tj T*" for line in lines)
                + " ET"
            ).encode()
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def file_part(name: str, raw: bytes) -> dict[str, Any]:
    return {
        "type": "file",
        "file": {"filename": name, "file_data": base64.b64encode(raw).decode()},
    }


class DocumentTests(unittest.TestCase):
    def test_pdf_all_sixty_pages_and_all_text_formats(self) -> None:
        result = extract_document(
            Attachment.model_validate(file_part("report.pdf", pdf())["file"])
        )
        self.assertEqual(result.pages, 60)
        for value in ("730 euros", "19 samples", "postpone launch", "[Page 60]"):
            self.assertIn(value, result.text)
        for extension in ("txt", "md", "csv"):
            text = "Beginning\nMiddle\nEnd\nUnicode: été"
            doc = extract_document(
                Attachment.model_validate(
                    file_part("data." + extension, text.encode())["file"]
                )
            )
            self.assertEqual(doc.text, text)

    def test_docx_includes_tables_in_order(self) -> None:
        document = Document()
        document.add_paragraph("Beginning")
        document.add_table(1, 1).cell(0, 0).text = "Middle"
        document.add_paragraph("End")
        buffer = io.BytesIO()
        document.save(buffer)
        result = extract_document(
            Attachment.model_validate(
                file_part("report.docx", buffer.getvalue())["file"]
            )
        )
        self.assertLess(result.text.index("Beginning"), result.text.index("Middle"))
        self.assertLess(result.text.index("Middle"), result.text.index("End"))

    def test_empty_scanned_encrypted_invalid_and_remote_rejected(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        buffer = io.BytesIO()
        writer.write(buffer)
        invalid = [("blank.pdf", buffer.getvalue(), "text layer")]
        writer.encrypt("test-only")
        buffer = io.BytesIO()
        writer.write(buffer)
        invalid.extend(
            [
                ("locked.pdf", buffer.getvalue(), "Password"),
                ("empty.txt", b"", "bytes"),
                ("bad.pdf", b"invalid", "extraction failed"),
                ("x.html", b"<p>text</p>", "Unsupported"),
            ]
        )
        for name, raw, message in invalid:
            with (
                self.subTest(name=name),
                self.assertRaisesRegex(
                    (DocumentError, ValueError), message if raw else ".*"
                ),
            ):
                extract_document(
                    Attachment.model_validate(file_part(name, raw)["file"])
                )
        with self.assertRaises(DocumentError):
            extract_document(
                Attachment(filename="remote.pdf", file_data="https://127.0.0.1/private")
            )

    def test_http_two_attachments_bypass_corpus_and_preserve_full_context(self) -> None:
        model = GatewayModel(
            "http://127.0.0.1:1",
            Configuration(
                input_eur_per_mtok=Decimal("0.1"),
                output_eur_per_mtok=Decimal("0.1"),
                max_tokens=1000,
            ),
        )

        async def generate(messages: list[dict[str, object]], timeout: float) -> Turn:
            self.assertEqual(model.tools, [])
            text = json.dumps(messages, ensure_ascii=False)
            for expected in (
                "730 euros",
                "19 samples",
                "postpone launch",
                "second attachment",
                "[Page 60]",
            ):
                self.assertIn(expected, text)
            return Turn(
                "Budget 730 euros; review 19 samples; postpone launch. Second attachment read."
            )

        with (
            tempfile.TemporaryDirectory() as root,
            patch.dict(os.environ, {"ATLAS_INTERACTION_DIR": root}),
            patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
            patch.object(model, "complete", generate),
            patch(
                "services.orchestrator.chat_pipeline.run",
                AsyncMock(side_effect=AssertionError("corpus route")),
            ) as corpus,
            TestClient(app) as client,
        ):
            reply = client.post(
                "/v1/chat/completions",
                json={
                    "messages": [
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": "Summarize both attachments"},
                                file_part("report.pdf", pdf()),
                                file_part("notes.txt", b"second attachment"),
                            ],
                        }
                    ]
                },
            )
            self.assertEqual(reply.status_code, 200, reply.text)
            self.assertIn(
                "postpone launch", reply.json()["choices"][0]["message"]["content"]
            )
            corpus.assert_not_called()
            row = json.loads(next(Path(root).glob("*.jsonl")).read_text())
            self.assertEqual(row["documents"][0]["pages"], 60)
            self.assertGreater(row["documents"][0]["extracted_characters"], 150000)
            self.assertFalse(row["documents"][0]["hierarchical_synthesis"])
            self.assertEqual(row["chunks_recuperes"], [])


class ActivityTests(unittest.IsolatedAsyncioTestCase):
    async def test_no_tool_status_persists_until_first_visible_token(self) -> None:
        ready = asyncio.Event()

        async def process(payload: ChatRequest, item: Interaction) -> None:
            await ready.wait()
            sink = sink_context.get()
            assert sink is not None
            await sink({"delta": {"content": "The complete answer.\n"}})
            item.state = "done"

        request = ChatRequest.model_validate(
            {
                "stream": True,
                "messages": [{"role": "user", "content": "Summarize this."}],
            }
        )
        with (
            tempfile.TemporaryDirectory() as root,
            patch.dict(os.environ, {"ATLAS_INTERACTION_DIR": root}),
        ):
            start = time.monotonic()
            stream = response(request, Interaction(), start, process).body_iterator
            assert isinstance(stream, AsyncGenerator)
            first = json.loads((await stream.__anext__())[6:])
            self.assertLess(time.monotonic() - start, 2)
            self.assertFalse(first["event"]["data"]["done"])
            previous = first
            for _ in range(30):
                second = json.loads((await stream.__anext__())[6:])
                elapsed = second["event"]["data"]["elapsed_seconds"]
                self.assertGreater(
                    elapsed, previous["event"]["data"]["elapsed_seconds"]
                )
                self.assertLessEqual(
                    elapsed - previous["event"]["data"]["elapsed_seconds"], 2
                )
                self.assertFalse(second["event"]["data"]["done"])
                previous = second
            self.assertGreaterEqual(time.monotonic() - start, 30)
            ready.set()
            events = [str(event) async for event in stream]
            content = next(
                json.loads(e[6:])
                for e in events
                if e.startswith("data: {")
                and json.loads(e[6:])
                .get("choices", [{}])[0]
                .get("delta", {})
                .get("content")
            )
            self.assertIn("complete answer", content["choices"][0]["delta"]["content"])
            self.assertTrue(content["event"]["data"]["done"])
            self.assertTrue(content["event"]["data"]["hidden"])

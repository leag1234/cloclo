"""Useful-content assertions on real authenticated M24 download recordings."""

import base64
import ast
import io
import re
from decimal import Decimal
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader
from PIL import Image
from packages.evidence import estimated_tokens
from services.orchestrator.documents import Attachment, extract_document


def spreadsheet_rows(sheet: ET.Element) -> list[list[str]]:
    """Compute supported formulas without trusting Excel's optional cached values."""
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cells = {cell.attrib["r"]: cell for cell in sheet.findall(".//s:c", ns)}

    def value(address: str, visiting: frozenset[str] = frozenset()) -> str:
        assert address not in visiting and len(visiting) < 100, "cyclic_formula"
        cell = cells[address]
        formula = cell.find("s:f", ns)
        if formula is None:
            return "".join(cell.itertext())
        match = re.fullmatch(
            r"SUM\(([A-Z]+)([0-9]+):([A-Z]+)([0-9]+)\)", formula.text or "", re.I
        )
        if match is None:
            expression = re.sub(r"(?<![<>=])=(?!=)", "==", formula.text or "")
            tree = ast.parse(expression, mode="eval")
            assert len(list(ast.walk(tree))) <= 100, "formula_too_large"

            def compute(node: ast.AST) -> Decimal | str | bool:
                if isinstance(node, ast.Constant) and type(node.value) in (
                    int,
                    float,
                    str,
                ):
                    return (
                        node.value
                        if isinstance(node.value, str)
                        else Decimal(str(node.value))
                    )
                if isinstance(node, ast.Name) and re.fullmatch(
                    r"[A-Z]+[0-9]+", node.id
                ):
                    return Decimal(value(node.id, visiting | {address}))
                if isinstance(node, ast.BinOp):
                    left, right = compute(node.left), compute(node.right)
                    assert isinstance(left, Decimal) and isinstance(right, Decimal)
                    if isinstance(node.op, ast.Add):
                        return left + right
                    if isinstance(node.op, ast.Sub):
                        return left - right
                    if isinstance(node.op, ast.Mult):
                        return left * right
                    if isinstance(node.op, ast.Div):
                        return left / right
                if (
                    isinstance(node, ast.Compare)
                    and len(node.ops) == 1
                    and isinstance(node.ops[0], ast.Eq)
                ):
                    return compute(node.left) == compute(node.comparators[0])
                if (
                    isinstance(node, ast.Call)
                    and isinstance(node.func, ast.Name)
                    and node.func.id.upper() == "IF"
                    and len(node.args) == 3
                    and not node.keywords
                ):
                    condition = compute(node.args[0])
                    assert isinstance(condition, bool)
                    return compute(node.args[1] if condition else node.args[2])
                raise AssertionError("unsupported_spreadsheet_formula")

            result = compute(tree.body)
            return format(result, "f") if isinstance(result, Decimal) else str(result)

        def column_number(label: str) -> int:
            number = 0
            for char in label.upper():
                number = number * 26 + ord(char) - ord("A") + 1
            return number

        left, right = column_number(match[1]), column_number(match[3])
        top, bottom = int(match[2]), int(match[4])
        assert left <= right and top <= bottom
        assert (right - left + 1) * (bottom - top + 1) <= 1000
        total = sum(
            (
                Decimal(value(key, visiting | {address}))
                for key in cells
                if (coordinate := re.fullmatch(r"([A-Z]+)([0-9]+)", key))
                and left <= column_number(coordinate[1]) <= right
                and top <= int(coordinate[2]) <= bottom
            ),
            Decimal(0),
        )
        return format(total, "f")

    return [
        [value(cell.attrib["r"]) for cell in row.findall("s:c", ns)]
        for row in sheet.findall(".//s:row", ns)
    ]


def validate(name: str, case: dict[str, Any], answer: dict[str, Any]) -> None:
    text = answer["choices"][0]["message"]["content"]
    tools = [
        row["request"]["call"]["name"]
        for row in case["exchanges"]
        if row["request"].get("kind") == "tool"
    ]
    assert "rag_search" not in tools, "documents_must_not_query_corpus"
    assert len(tools) <= 10
    assert case["response"]["seconds"] <= 120
    attached = any(
        isinstance(m["content"], list)
        and any(p["type"] == "file" for p in m["content"])
        for m in case["request"]["messages"]
    )
    # Owner ruling 2026-09-21: file production shares the attachment ceiling.
    produces_files = bool(case["artifacts"])
    assert (
        0
        < answer["atlas"]["cost_eur"]
        <= (0.30 if attached or produces_files else 0.10)
    )
    artifacts = {
        artifact["url"].rsplit(".", 1)[-1]: base64.b64decode(
            artifact["bytes"], validate=True
        )
        for artifact in case["artifacts"]
    }
    for artifact in case["artifacts"]:
        assert artifact["status"] == 200
        # IDs are substituted during replay; the filename and authenticated route remain.
        filename = artifact["url"].rsplit("/", 1)[-1]
        assert filename in text and "/api/v1/terminals/atlas-files/files/serve/" in text
    if name == "J42":
        assert set(artifacts) == {"pptx", "pdf"}
        with ZipFile(io.BytesIO(artifacts["pptx"])) as archive:
            slides = [
                n
                for n in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
            ]
            assert len(slides) >= 4
            contents = [
                " ".join(ET.fromstring(archive.read(n)).itertext()) for n in slides
            ]
            assert len(set(contents)) >= 4 and all(
                content.strip() for content in contents
            )
            assert "murena" in " ".join(contents).lower()
        pdf = PdfReader(io.BytesIO(artifacts["pdf"]))
        assert len(pdf.pages) == len(slides)
        assert "murena" in " ".join(page.extract_text() for page in pdf.pages).lower()
    elif name == "J43":
        document = Document(io.BytesIO(artifacts["docx"]))
        paragraphs = " ".join(p.text for p in document.paragraphs)
        assert "employés travaillent" in paragraphs and "outils adaptés" in paragraphs
        assert (
            "avons terminé" in paragraphs
            and "résultats sont satisfaisants" in paragraphs
        )
        assert (
            sum(
                p.style is not None and p.style.name.startswith("Heading")
                for p in document.paragraphs
            )
            >= 2
        )
        with ZipFile(io.BytesIO(artifacts["docx"])) as archive:
            assert b"TOC" in archive.read("word/document.xml")
        assert re.search(r"régénéré|regenerated|en place|in place", text, re.I)
    elif name == "J44":
        parts = case["request"]["messages"][0]["content"]
        original = base64.b64decode(
            next(p["file"]["file_data"] for p in parts if p["type"] == "file")
        )
        with (
            ZipFile(io.BytesIO(artifacts["xlsx"])) as archive,
            ZipFile(io.BytesIO(original)) as before,
        ):
            assert archive.read("xl/worksheets/sheet1.xml") == before.read(
                "xl/worksheets/sheet1.xml"
            ), "original_sheet_changed"
            charts = [
                n
                for n in archive.namelist()
                if re.fullmatch(r"xl/charts/chart\d+\.xml", n)
            ]
            assert charts and "Comparaison" in archive.read(charts[0]).decode()
            sheet = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
            rows = spreadsheet_rows(sheet)
            total = next(row for row in rows if "Total" in row)
            assert "200" in total and "260" in total and "60" in total, (
                "incorrect_march_totals"
            )
    elif name in {"J45", "J47"}:
        assert "730" in text and "19" in text
        assert re.search(r"report|postpon", text, re.I)
        if name == "J45":
            assert "note.txt" in text and "60" in text
        else:
            parts = case["request"]["messages"][0]["content"]
            attachment = next(p["file"] for p in parts if p["type"] == "file")
            extracted = extract_document(Attachment.model_validate(attachment))
            assert estimated_tokens(extracted.text) >= 175000
            assert str(extracted.pages) in text
    elif name == "J46":
        assert re.search(r"mot de passe|password|chiffr", text, re.I)
        assert re.search(
            r"ne peux (?:donc )?pas|impossible|cannot|illisible", text, re.I
        )
    elif name == "J48":
        assert "web_search" in tools
        assert re.search(r"https?://[^\s)]*(?:cpcwiki|cpctech)[^\s)]*", text, re.I), (
            "missing_cpc_citation"
        )
        assert re.search(r"consult|2026-09-18", text, re.I)
        # Measured CPC timings: https://www.cpctech.cpcwiki.de/docs/instrtim.html
        assert re.search(r"OUTI[^\n]{0,100}\b5\s*(?:[µμu]s|microsecond)", text, re.I)
        assert re.search(r"OTIR[^\n]{0,120}\b6\s*(?:[µμu]s|microsecond)", text, re.I)
    elif name == "formats_read":
        rows = [
            [
                cell.strip().strip("`*").casefold()
                for cell in line.strip().strip("|").split("|")
            ]
            for line in text.splitlines()
            if line.strip().startswith("|")
        ]
        headers = next(
            (
                row
                for row in rows
                if {"fichier", "article"} <= set(row)
                and any(label in row for label in ("quantité", "quantite"))
            ),
            None,
        )
        assert headers is not None, "missing_comparison_headers"
        columns = [
            headers.index("fichier"),
            headers.index("article"),
            headers.index("quantité" if "quantité" in headers else "quantite"),
        ]
        for filename, (article, count) in {
            "lettre.odt": ("aubergine", 37),
            "inventaire.ods": ("abricot", 43),
            "diapos.odp": ("cerise", 59),
            "diapos.pptx": ("cerise", 59),
            "stock.csv": ("datte", 61),
            "notes.md": ("figue", 71),
            "note.txt": ("groseille", 83),
        }.items():
            assert re.search(
                re.escape(filename) + r"[^\n]*\b" + str(count) + r"\b", text
            ), "missing_format_content"
            assert any(
                len(row) == len(headers)
                and [row[column] for column in columns]
                == [filename, article, str(count)]
                for row in rows
            ), "missing_comparison_row"
    elif name == "formats_write":
        assert set(artifacts) == {"csv", "md", "png"}
        for extension in ("csv", "md"):
            content = artifacts[extension].decode("utf-8-sig")
            assert all(str(number) in content for number in (10, 15, 25))
        with Image.open(io.BytesIO(artifacts["png"])) as graph:
            assert graph.format == "PNG" and graph.width >= 320 and graph.height >= 200
            assert len(set(graph.convert("RGB").getdata())) > 100
    else:
        raise AssertionError("unknown_m24_journey")

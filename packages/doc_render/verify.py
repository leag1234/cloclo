"""Reopen files and distinguish physical pagination from package checks."""

from dataclasses import dataclass
from pathlib import Path
from typing import Literal, cast

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pypdf import PdfReader

from .schema import Content


class VerificationError(ValueError):
    """An artifact cannot be published under the requested verification policy."""


@dataclass(frozen=True)
class Verification:
    unit: Literal["page", "slide", "document", "worksheet"]
    count: int
    text: tuple[str, ...]
    pagination_verified: bool


def verify(
    path: Path,
    *,
    expected: Content | None = None,
    expected_count: int | None = None,
    require_pagination: bool = False,
) -> Verification:
    """Parsing errors propagate; blank units and missing content fail explicitly."""
    text: list[str] = []
    unit: Literal["page", "slide", "document", "worksheet"]
    if path.suffix == ".pdf":
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise VerificationError("encrypted_pdf")
        text = [page.extract_text() for page in reader.pages]
        unit = "page"
    elif path.suffix == ".pptx":
        deck = Presentation(str(path))
        for slide in deck.slides:
            parts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(cast(Shape, shape).text)
                if shape.has_table:
                    parts.extend(
                        cell.text
                        for row in cast(GraphicFrame, shape).table.rows
                        for cell in row.cells
                    )
            text.append("\n".join(parts))
        unit = "slide"
    elif path.suffix == ".docx":
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        text = ["\n".join(parts)]
        unit = "document"
    elif path.suffix == ".xlsx":
        workbook = load_workbook(path, data_only=False, read_only=True)
        try:
            text = [
                "\n".join(
                    str(cell.value)
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.value is not None
                )
                for sheet in workbook.worksheets
            ]
        finally:
            workbook.close()
        unit = "worksheet"
    else:
        raise VerificationError("unsupported_output_format")
    if not text or any(not value.strip() for value in text):
        raise VerificationError("empty_document_or_unit")
    if expected_count is not None and len(text) != expected_count:
        raise VerificationError(f"unit_count_mismatch:{len(text)}:{expected_count}")
    paginated = unit in {"page", "slide"}
    if require_pagination and not paginated:
        raise VerificationError("office_layout_engine_required")
    if expected:
        units = ["".join(value.split()) for value in text]
        if paginated:
            seen: set[str] = set()
            headings = sorted(
                {"".join(section.heading.split()) for section in expected.sections},
                key=len,
                reverse=True,
            )
            for index, value in enumerate(units):
                for heading in headings:
                    if value.startswith(heading):
                        if heading in seen:
                            units[index] = value[len(heading) :]
                        break
                seen.update(heading for heading in headings if heading in value)
        combined = "".join(units)
        if any("".join(value.split()) not in combined for value in expected.texts()):
            raise VerificationError("missing_expected_text")
    return Verification(unit, len(text), tuple(text), paginated)

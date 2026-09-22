"""Bounded local extraction for POC-F2; documents never trigger network I/O."""

import argparse
import json
import logging
from html.parser import HTMLParser
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PyPdfError

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 1_000_000
LOG = logging.getLogger(__name__)


class InvalidDocument(ValueError):
    """Unsupported, unsafe, empty or malformed document."""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        LOG.warning(message)


class HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "template"}:
            self.hidden += 1
        if not self.hidden and tag in {"p", "br", "div", "li", "h1", "h2", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "template"}:
            self.hidden = max(0, self.hidden - 1)
        if not self.hidden and tag in {"p", "div", "li", "h1", "h2", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.parts.append(data)


def extract(path: Path) -> str:
    """Extract a regular local file; the caller supplies corpus ownership metadata."""
    try:
        if path.is_symlink() or not path.is_file():
            raise InvalidDocument("invalid_file")
        if path.stat().st_size > MAX_BYTES:
            raise InvalidDocument(
                f"invalid_file: {path.stat().st_size} bytes; limit {MAX_BYTES} bytes"
            )
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            reader = PdfReader(path, strict=True)
            if reader.is_encrypted:
                raise InvalidDocument("encrypted_pdf")
            parts: list[str] = []
            total = 0
            for page in reader.pages:
                value = page.extract_text()
                total += len(value)
                if total > MAX_TEXT:
                    raise InvalidDocument(
                        f"text_limit: {total} characters; limit {MAX_TEXT} characters"
                    )
                parts.append(value)
            text = "\n".join(parts)
        elif suffix == ".docx":
            with ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_TEXT:
                    raise InvalidDocument(
                        f"archive_limit: {sum(i.file_size for i in archive.infolist())} bytes; limit {MAX_TEXT} bytes"
                    )
            doc = Document(str(path))
            text = "\n".join(
                "\n".join("\t".join(c.text for c in row.cells) for row in block.rows)
                if isinstance(block, Table)
                else block.text
                for block in doc.iter_inner_content()
            )
        elif suffix in {".md", ".html"}:
            text = path.read_text(encoding="utf-8")
            if suffix == ".html":
                parser = HTMLText()
                parser.feed(text)
                parser.close()
                text = "\n".join(
                    line.strip()
                    for line in "".join(parser.parts).splitlines()
                    if line.strip()
                )
        else:
            raise InvalidDocument("unsupported_format")
        if len(text) > MAX_TEXT:
            raise InvalidDocument(
                f"text_limit: {len(text)} characters; limit {MAX_TEXT} characters"
            )
        if not text.strip():
            raise InvalidDocument("empty_or_large_text")
    except (
        OSError,
        UnicodeError,
        BadZipFile,
        PyPdfError,
        XMLSyntaxError,
        KeyError,
    ) as exc:
        LOG.warning('{"event":"extraction_failed","code":"invalid_document"}')
        raise InvalidDocument("invalid_document") from exc
    LOG.info(
        json.dumps({"event": "extracted", "format": suffix, "characters": len(text)})
    )
    return text.strip()


def split_text(text: str) -> list[str]:
    """Version 1: fixed character windows with overlap, preserving exact passages."""
    return [
        text[start : start + 2400].strip()
        for start in range(0, len(text), 2200)
        if text[start : start + 2400].strip()
    ]


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extract a local document into text chunks"
    )
    parser.add_argument("path", type=Path)
    args = parser.parse_args()
    # Binary parsers also need CPU/address-space limits before reading untrusted input.
    import resource

    resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
    logging.basicConfig(level=logging.INFO)
    print(json.dumps(split_text(extract(args.path)), ensure_ascii=False))


if __name__ == "__main__":
    main()

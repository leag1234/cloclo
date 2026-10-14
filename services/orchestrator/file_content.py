"""Render only delivered answer text; never execute model-authored fallback code."""

import json
import re
import shlex
import textwrap
import unicodedata

from packages.file_answers import strip_file_promises


def folded(text: str) -> str:
    return "".join(
        c
        for c in unicodedata.normalize("NFKD", text.casefold())
        if not unicodedata.combining(c)
    )


def failure_notice(text: str) -> bool:
    value = folded(text)
    return bool(
        re.search(r"file|fichier|document|datei|archivo", value)
        and re.search(
            r"pas pu|impossible|echou|could not|cannot|failed|unable|keine|nicht|no se pudo|ningun|non.*pubblic|nessun",
            value,
        )
    )


def answer_content(text: str, *, preserve_code: bool = False) -> str:
    text = text if preserve_code else re.sub(r"```.*?(?:```|$)", "", text, flags=re.S)
    if not preserve_code:
        text = strip_file_promises(text)
    lines = []
    for line in text.splitlines():
        value = folded(line)
        if failure_notice(line) and not re.search(
            r"sandbox:|https?://|untrusted_tool_history", line
        ):
            lines.append(line)
            continue
        # Publication comes exclusively from the terminal's verified file list.
        if (
            re.search(
                r"(?:j'ai|i have|i've|ich habe|he |ho ).{0,40}(?:cree|created|generated|gener|erstell|creado|creato).{0,60}(?:presentation|file|fichier|document|pdf|pptx)",
                value,
            )
            or "untrusted_tool_history" in line
            or "sandbox:" in line
            or re.search(r"download|telecharg|herunterlad|descarg|scaric", value)
            or re.search(
                r"(?:je vais|je publie|i will|i'll|ich werde|voy a|procedero).*?(?:fichier|file|document|pptx|pdf|datei)",
                value,
            )
            or re.search(
                r"(?:fichier|file|document|pptx|pdf|datei).*(?:genere|generated|cree|created|publie|published|disponible|available)",
                value,
            )
            or re.search(
                r"^\s*#+\s*(?:generation|creation|download|telecharg).*?(?:document|file|fichier)",
                value,
            )
        ):
            continue
        lines.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def requested_formats(question: str) -> tuple[str, ...]:
    formats = tuple(
        ext
        for ext in ("pptx", "pdf", "docx", "xlsx", "csv", "md")
        if re.search(r"\b" + ext + r"\b", question, re.I)
    )
    if not formats and re.search(
        r"presentation|présentation|slides|folien|diapositiv", question, re.I
    ):
        return ("pptx",)
    return formats


def sections(text: str) -> list[tuple[str, list[str]]]:
    """Split at the answer's headings, then wrap its own paragraphs and bullets."""
    groups: list[tuple[str, list[str]]] = []
    title = ""
    body: list[str] = []
    for line in answer_content(text).splitlines():
        line = line.strip()
        if failure_notice(line):
            continue
        heading = re.fullmatch(r"(?:#{1,6}\s+(.+)|\*\*(\d+[.)].+?)\*\*)", line)
        if heading:
            if body:
                groups.append((title, body))
            title = (heading[1] or heading[2]).replace("**", "")
            body = []
        elif line and not re.fullmatch(r"[-*_]{3,}", line):
            clean = re.sub(r"^\s*[*+-]\s+", "• ", line).replace("**", "")
            body.extend(
                textwrap.wrap(
                    clean, width=82, break_long_words=True, break_on_hyphens=False
                )
            )
    if body:
        groups.append((title, body))
    # Each slide has real body text. No minimum count and no padding.
    return [
        (title, lines[start : start + 7])
        for title, lines in groups
        for start in range(0, len(lines), 7)
    ]


def render_command(text: str, formats: tuple[str, ...]) -> tuple[str, list[str]]:
    selected = [ext for ext in formats if ext in {"pptx", "pdf"}]
    slides = sections(text)
    if not selected or not slides:
        return "", []
    data = json.dumps(slides, ensure_ascii=False)
    code = """import json, subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
slides = json.loads(DATA)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
for title, lines in slides:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if title:
        box = slide.shapes.add_textbox(Inches(.6), Inches(.3), Inches(12.1), Inches(1.2))
        box.text_frame.word_wrap = True
        box.text_frame.text = title
        for paragraph in box.text_frame.paragraphs:
            paragraph.font.size = Pt(24)
    box = slide.shapes.add_textbox(Inches(.6), Inches(1.7), Inches(12.1), Inches(5.3))
    box.text_frame.word_wrap = True
    box.text_frame.text = "\\n".join(lines)
    for paragraph in box.text_frame.paragraphs:
        paragraph.font.size = Pt(20)
prs.save("answer.pptx")
""".replace("DATA", repr(data))
    if "pdf" in selected:
        code += """subprocess.run(["libreoffice", "-env:UserInstallation=file:///tmp/atlas-render-" + Path.cwd().name, "--headless", "--convert-to", "pdf", "--outdir", ".", "answer.pptx"], check=True, timeout=25)
assert Path("answer.pdf").stat().st_size > 0
"""
    return "python3 -c " + shlex.quote(code), ["answer." + ext for ext in selected]


def presentable(
    text: str, formats: tuple[str, ...], *, structured: bool = False
) -> bool:
    """A content turn needs three substantive titled sections for slides, one otherwise.

    Existing prose can still be paginated. New content must have explicit section
    titles; page wrapping cannot turn a single paragraph into accepted sections.
    Plans, apologies and promised generation never count as deliverable content.
    """
    value = folded(answer_content(text))
    if re.search(
        r"plan (?:de la|de|for the) (?:presentation|document)|"
        r"(?:je vais|i will|i'll|ich werde|voy a|procedero a).{0,100}"
        r"integrer|(?:je vais|i will|i'll|ich werde|voy a|procedero a).{0,100}"
        r"(?:gener|creer|create|prepare|publish|publier|erstell|crear)|"
        r"(?:sorry|desole|apolog).{0,60}(?:cannot|impossible|pas)",
        value,
    ):
        return False
    groups = sections(text)
    minimum = 3 if "pptx" in formats else 1
    if structured:
        titles = {
            title for title, lines in groups if title and len(" ".join(lines)) >= 60
        }
        if len(titles) < minimum:
            return False
    return (
        len(groups) >= minimum
        and sum(len(" ".join(lines)) for _, lines in groups) >= 100 * minimum
    )
